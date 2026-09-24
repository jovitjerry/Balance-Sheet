from __future__ import annotations
import io
import logging
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from app.core.config import Settings
from app.core.errors import UnreadableDocumentError
from app.core.schemas import ParserKind, PositionedWord, PreliminaryExtraction, RawTable, SourceFile, SourcePage
from app.modules.ingestion.ocr import OcrEngine, OcrUnavailable, render_pdf_page
from app.modules.ingestion.validation import UploadKind, looks_like_pdf
logger = logging.getLogger(__name__)
MAX_PAGES = 200

async def parse(data: bytes, source: SourceFile, *, kind: UploadKind | None=None, settings: Settings | None=None, ocr: OcrEngine | None=None) -> PreliminaryExtraction:
    from app.core.config import get_settings
    config = settings or get_settings()
    resolved = kind or (UploadKind.PDF if looks_like_pdf(data) else UploadKind.XLSX)
    if resolved is UploadKind.PDF:
        return _parse_pdf(data, config, ocr)
    return _parse_xlsx(data)

def _parse_pdf(data: bytes, settings: Settings, ocr: OcrEngine | None) -> PreliminaryExtraction:
    import pdfplumber
    pages: list[SourcePage] = []
    warnings: list[str] = []
    try:
        pdf = pdfplumber.open(io.BytesIO(data))
    except Exception as exc:
        logger.warning('PDF could not be opened: %s', exc)
        raise UnreadableDocumentError('The PDF could not be opened. It may be corrupt, or password protected.') from exc
    try:
        if not pdf.pages:
            raise UnreadableDocumentError('The PDF contains no pages.')
        if len(pdf.pages) > MAX_PAGES:
            raise UnreadableDocumentError(f'The PDF has more than {MAX_PAGES} pages, which is beyond what this system processes.')
        for index, page in enumerate(pdf.pages):
            pages.append(_parse_pdf_page(data, page, index, settings, ocr, warnings))
    finally:
        pdf.close()
    if not any((page.text.strip() for page in pages)):
        raise UnreadableDocumentError('No readable text could be extracted from the document.')
    ocr_pages = sum((1 for page in pages if page.ocr_used))
    if ocr_pages == 0:
        parser = ParserKind.PDF_TEXT
    elif ocr_pages == len(pages):
        parser = ParserKind.PDF_OCR
    else:
        parser = ParserKind.PDF_MIXED
    return PreliminaryExtraction(parser=parser, pages=pages, warnings=warnings)

def _parse_pdf_page(data: bytes, page: object, index: int, settings: Settings, ocr: OcrEngine | None, warnings: list[str]) -> SourcePage:
    text = (page.extract_text() or '').strip()
    char_count = len(page.chars)
    if char_count >= settings.ocr_min_chars_per_page and text:
        return SourcePage(index=index, text=text, tables=_extract_pdf_tables(page, warnings, index), ocr_used=False, words=_extract_pdf_words(page, index, warnings), page_width=float(page.width), page_height=float(page.height))
    return _ocr_pdf_page(data, index, settings, ocr, warnings)

def _ocr_pdf_page(data: bytes, index: int, settings: Settings, ocr: OcrEngine | None, warnings: list[str]) -> SourcePage:
    if not settings.ocr_enabled:
        raise OcrUnavailable(f'Page {index + 1} is a scanned image and OCR is disabled on this server, so it cannot be read.')
    if ocr is None or not ocr.available():
        raise OcrUnavailable(f'Page {index + 1} is a scanned image and requires OCR, which is not available on this server.')
    try:
        image = render_pdf_page(data, index, dpi=settings.ocr_dpi)
    except OcrUnavailable:
        raise
    except Exception as exc:
        logger.warning('Could not rasterise page %d: %s', index, exc)
        raise UnreadableDocumentError(f'Page {index + 1} could not be rendered for reading.') from exc
    result = ocr.recognise(image)
    if not result.words:
        warnings.append(f'Page {index + 1}: OCR found no legible text.')
    return SourcePage(index=index, text=result.text, tables=[], ocr_used=True, ocr_confidence=result.confidence, words=result.words, page_width=float(result.width), page_height=float(result.height))

def _extract_pdf_words(page: object, index: int, warnings: list[str]) -> list[PositionedWord]:
    try:
        found = page.extract_words()
    except Exception as exc:
        logger.info('Word extraction failed on page %d: %s', index, exc)
        warnings.append(f'Page {index + 1}: word positions could not be extracted.')
        return []
    return [PositionedWord(text=word['text'], left=float(word['x0']), top=float(word['top']), width=max(0.0, float(word['x1']) - float(word['x0'])), height=max(0.0, float(word['bottom']) - float(word['top']))) for word in found if word.get('text', '').strip()]

def _extract_pdf_tables(page: object, warnings: list[str], index: int) -> list[RawTable]:
    try:
        found = page.extract_tables()
    except Exception as exc:
        logger.info('Table extraction failed on page %d: %s', index, exc)
        warnings.append(f'Page {index + 1}: tables could not be extracted.')
        return []
    tables: list[RawTable] = []
    for table in found:
        rows = [[None if cell is None else str(cell).strip() for cell in row] for row in table]
        if any((any((cell for cell in row)) for row in rows)):
            tables.append(RawTable(rows=rows))
    return tables

def _parse_xlsx(data: bytes) -> PreliminaryExtraction:
    from openpyxl import load_workbook
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        logger.warning('Workbook could not be opened: %s', exc)
        raise UnreadableDocumentError('The Excel workbook could not be read. The file appears to be corrupt or incomplete.') from exc
    pages: list[SourcePage] = []
    warnings: list[str] = []
    try:
        for index, sheet in enumerate(workbook.worksheets):
            rows = [[_cell_to_str(value) for value in row] for row in sheet.iter_rows(values_only=True)]
            rows = [row for row in rows if any((cell for cell in row))]
            pages.append(SourcePage(index=index, name=sheet.title, text='\n'.join((' '.join((cell for cell in row if cell)) for row in rows)), tables=[RawTable(rows=rows)] if rows else [], ocr_used=False))
    finally:
        workbook.close()
    if not pages or not any((page.text.strip() for page in pages)):
        raise UnreadableDocumentError('The workbook is empty - no readable cells were found.')
    return PreliminaryExtraction(parser=ParserKind.EXCEL, pages=pages, warnings=warnings)

def _cell_to_str(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value, str):
        return value.strip() or None
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, int):
        return str(value)
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, float):
        try:
            decimal = Decimal(str(value))
        except InvalidOperation:
            return str(value)
        if decimal == decimal.to_integral_value():
            return str(decimal.quantize(Decimal(1)))
        return str(decimal.normalize())
    return str(value)
__all__ = ['MAX_PAGES', 'parse']
